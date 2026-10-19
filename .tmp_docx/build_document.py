from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[1]
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = section.bottom_margin = Inches(.7)
section.left_margin = section.right_margin = Inches(.8)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(24)
doc.styles['Heading 1'].font.size = Pt(15)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(13)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(6)
doc.core_properties.title = 'Cancelación de complementos de pago en facturas PPD'
doc.core_properties.subject = 'Documentación del cambio en Entintados PDV para Odoo 19'
doc.core_properties.author = 'Entintados PDV'

def p(text, style=None):
    return doc.add_paragraph(text, style)

def h(text):
    doc.add_heading(text, level=1)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for cell, text, width in zip(t.rows[0].cells, headers, widths):
        cell.width = Inches(width)
        cell.text = text
    for row in rows:
        for cell, text, width in zip(t.add_row().cells, row, widths):
            cell.width = Inches(width)
            cell.text = text
    for i, row in enumerate(t.rows):
        for cell in row.cells:
            props = cell._tc.get_or_add_tcPr()
            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), 'E7EDF3' if i == 0 else 'FFFFFF')
            props.append(shading)
            borders = OxmlElement('w:tcBorders')
            for side in ['top', 'left', 'bottom', 'right']:
                border = OxmlElement('w:' + side)
                for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                    border.set(qn('w:' + key), value)
                borders.append(border)
            props.append(borders)
            margins = OxmlElement('w:tcMar')
            for side in ['top', 'left', 'bottom', 'right']:
                node = OxmlElement('w:' + side)
                node.set(qn('w:w'), '100')
                node.set(qn('w:type'), 'dxa')
                margins.append(node)
            props.append(margins)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.bold = i == 0
    return t

p('Cancelación de complementos de pago en facturas PPD', 'Title')
p('Entintados PDV · Odoo 19 · 13 de octubre de 2026', 'Subtitle')
p('Se incorporó la cancelación de los complementos de pago relacionados antes de solicitar la cancelación del timbre fiscal de una factura PPD. Si un complemento no puede cancelarse, se detiene la cancelación de la factura y se registra el error para permitir un reintento.')
h('Objetivo y alcance')
p('El cambio aplica a facturas de cliente publicadas con estado fiscal enviado y método de pago PPD. Se ejecuta al solicitar la cancelación desde Odoo mediante el flujo nativo de CFDI, incluidos los reintentos. Las facturas PUE y los documentos que no cumplen estas condiciones siguen el comportamiento nativo.')
p('La búsqueda utiliza el historial fiscal de documentos relacionados con la factura y la misma compañía. Esto permite encontrar complementos timbrados aunque el pago ya no esté conciliado contablemente con la factura.')
h('Flujo de cancelación')
for text in [
    '1. Verificar que la factura cumpla las condiciones del alcance.',
    '2. Buscar sus documentos CFDI de pago enviados y bloquear los registros involucrados mediante el mecanismo nativo de Odoo.',
    '3. Omitir los complementos que ya tengan un documento de cancelación registrado.',
    '4. Cancelar cada complemento con el método nativo de Odoo. Este determina el motivo según exista o no un documento de sustitución.',
    '5. Si falla un complemento, registrar el error en los documentos CFDI de la factura y detener el flujo. Si todos se cancelan, continuar con la cancelación nativa de la factura y el motivo elegido por el usuario.',
]:
    p(text)
h('Efectos operativos')
p('Si un complemento incluye varias facturas, se cancela el CFDI de pago completo. La cancelación no puede limitarse a la referencia de una sola factura dentro del mismo complemento.')
p('Cuando un complemento falla después de que otros se cancelaron correctamente, se conserva el historial de las cancelaciones exitosas. En el siguiente intento se omiten esos documentos y se procesa lo pendiente.')
p('La aprobación y los estados fiscales finales dependen del flujo nativo de Odoo y de las respuestas del PAC y del SAT. Este cambio intercepta la solicitud desde Odoo; no añade un flujo específico para cancelaciones iniciadas directamente en el portal del SAT.')

doc.add_page_break()
h('Archivos y cambios técnicos')
table(['Archivo', 'Cambio'], [
    ('models/account_move.py', 'Se extendió _l10n_mx_edi_cfdi_invoice_try_cancel para cancelar los complementos antes de continuar con la factura.'),
    ('__manifest__.py', 'Se agregó l10n_mx_edi como dependencia explícita del módulo.'),
    ('tests/test_cfdi_payment_cancel.py', 'Se incorporaron cinco pruebas aisladas del flujo con ORM y llamadas fiscales simulados.'),
], [2.4, 4.5])
p('El método nativo _l10n_mx_edi_cfdi_invoice_try_cancel_payment realiza las cancelaciones de pago. La comprobación de documentos de cancelación evita volver a cancelar CFDI que permanecen en el historial como enviados.')
p('Ante un fallo se utiliza _l10n_mx_edi_cfdi_invoice_document_cancel_failed para dejar constancia en la factura. El flujo retorna sin lanzar una excepción deliberada que revierta el historial de las llamadas al PAC que ya tuvieron éxito.')
h('Pruebas realizadas')
table(['Caso', 'Resultado'], [
    ('Varios complementos', 'Se procesan todos antes de cancelar la factura.'),
    ('Fallo de un complemento', 'La factura no se cancela y se registra el error.'),
    ('Reintento', 'Se omite el complemento previamente cancelado.'),
    ('Factura sin complementos', 'Continúa la cancelación nativa.'),
    ('Otros flujos', 'PUE, borradores, CFDI cancelados y facturas de proveedor no activan la cancelación de complementos.'),
], [2.4, 4.5])
p('Resultado observado: 5 pruebas aprobadas dentro del contenedor odoo_o19. Las pruebas simulan el PAC y el ORM; no cancelan documentos fiscales reales ni sustituyen una validación integral en una base de pruebas de Odoo.')
p('Comando utilizado: docker exec odoo_o19 python3 /mnt/extra-addons/modules/entintados_pdv/tests/test_cfdi_payment_cancel.py')
h('Activación y validación funcional')
p('Para activar el cambio, actualizar el módulo entintados_pdv en la base correspondiente y reiniciar el servicio de Odoo. Estas acciones no se ejecutaron como parte de la implementación documentada.')
p('Antes de usarlo en operación, validar en una base de pruebas con PAC de prueba: una factura PPD con varios pagos, un complemento que incluya varias facturas y una respuesta de error del PAC. Confirmar los estados de los documentos, el registro del error y el comportamiento del reintento.')
h('Uso por el usuario')
p('Desde la factura PPD, abrir el flujo habitual de cancelación CFDI y seleccionar el motivo de cancelación. Si se informa un fallo de un complemento, revisar sus documentos CFDI, resolver la causa y reintentar la cancelación de la factura.')
out = root / 'docs' / 'Cambios_cancelacion_CFDI_facturas_PPD.docx'
doc.save(out)
print(out)
